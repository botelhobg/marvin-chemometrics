from __future__ import annotations
import csv
import io
import json
import zipfile
from pathlib import Path
import numpy as np
import pandas as pd

VERSION = '0.5.0'
MAX_BYTES = 20 * 1024**2
MAX_EXPANDED = 64 * 1024**2
MAX_CELLS = 2_000_000
ROLES = ['X', 'Y', 'Identificação', 'Classe', 'Grupo', 'Ignorar']


def check_zip(raw, exact=None):
    if len(raw) > MAX_BYTES:
        raise ValueError('Arquivo maior que 20 MB.')
    z = zipfile.ZipFile(io.BytesIO(raw))
    info = z.infolist()
    names = [i.filename for i in info]
    if len(names) > 2000 or len(names) != len(set(names)) or sum(i.file_size for i in info) > MAX_EXPANDED:
        z.close()
        raise ValueError('Pacote excede limites ou contém entradas duplicadas.')
    if any(i.flag_bits & 1 for i in info) or (exact is not None and set(names) != set(exact)):
        z.close()
        raise ValueError('Estrutura de pacote não reconhecida.')
    return z


def normalize_table(df):
    if df.empty or df.shape[0] > 20_000 or df.shape[1] > 10_000 or df.size > MAX_CELLS:
        raise ValueError('Tabela vazia ou acima dos limites: 20 mil amostras, 10 mil colunas, 2 milhões de células.')
    columns = [str(c).strip() for c in df.columns]
    if any(not c for c in columns) or len(columns) != len(set(columns)):
        raise ValueError('Os nomes das colunas devem ser preenchidos e únicos.')
    df = df.copy()
    df.columns = columns
    return df.fillna('').astype(str).reset_index(drop=True)


def excel_sheets(raw):
    with check_zip(raw):
        pass
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def read_table(raw, name, sep=';', decimal=',', encoding='utf-8-sig', header=True, sheet=None):
    if len(raw) > MAX_BYTES:
        raise ValueError('Arquivo maior que 20 MB.')
    suffix = Path(name).suffix.lower()
    if suffix in ['.csv', '.tsv', '.txt']:
        text = raw.decode(encoding)
        rows = []
        for row in csv.reader(io.StringIO(text), delimiter=sep):
            if not row or all(not x.strip() for x in row):
                continue
            rows.append(row)
            if len(rows) > 20_001 or len(row) > 10_000 or len(rows)*len(row) > MAX_CELLS + 10_000:
                raise ValueError('Tabela excede os limites de tamanho.')
        if not rows or any(len(r) != len(rows[0]) for r in rows):
            raise ValueError('Linhas com larguras diferentes ou arquivo vazio. Confira o separador.')
        cols = rows.pop(0) if header else [f'V{i+1}' for i in range(len(rows[0]))]
        df = pd.DataFrame(rows, columns=cols)
    elif suffix == '.xlsx':
        with check_zip(raw):
            pass
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            ws = wb[sheet or wb.sheetnames[0]]
            if ws.max_row > 20_001 or ws.max_column > 10_000 or ws.max_row * ws.max_column > MAX_CELLS + 10_000:
                raise ValueError('Planilha excede os limites. Remova linhas/colunas vazias formatadas.')
            rows = list(ws.values)
        finally:
            wb.close()
        if not rows:
            raise ValueError('Planilha vazia.')
        cols = rows.pop(0) if header else [f'V{i+1}' for i in range(len(rows[0]))]
        if any(c is None for c in cols):
            raise ValueError('Cabeçalhos vazios na planilha.')
        # Excel stores numeric cells numerically, independently of display locale.
        df = pd.DataFrame(rows, columns=cols).dropna(how='all')
        decimal = '.'
    else:
        raise ValueError('Use CSV, TSV, TXT ou XLSX.')
    return normalize_table(df), decimal


def numeric_column(series, decimal):
    s = series.astype(str).str.strip()
    if decimal == ',':
        s = s.str.replace(',', '.', regex=False)
    return pd.to_numeric(s.replace('', np.nan), errors='coerce')


def suggest_roles(df, decimal):
    result = {}
    for col in df:
        nonempty = df[col].str.strip().ne('')
        nums = numeric_column(df[col], decimal)
        result[col] = 'X' if nonempty.any() and nums[nonempty].notna().all() else 'Ignorar'
    return result


def default_config(df, decimal='.'):
    return {'roles': suggest_roles(df, decimal), 'decimal': decimal, 'excluded': [],
            'axis_mode': 'Posição', 'axis_label': 'Variável', 'axis_unit': '', 'notes': ''}


def validate_config(df, cfg):
    if not isinstance(cfg, dict) or set(cfg) != {'roles','decimal','excluded','axis_mode','axis_label','axis_unit','notes'}:
        raise ValueError('Configuração de projeto inválida.')
    roles = cfg['roles']
    if not isinstance(roles, dict) or set(roles) != set(df.columns) or any(v not in ROLES for v in roles.values()):
        raise ValueError('Papéis das colunas inválidos.')
    if cfg['decimal'] not in ['.', ','] or cfg['axis_mode'] not in ['Posição', 'Cabeçalhos numéricos']:
        raise ValueError('Formato numérico ou eixo inválido.')
    for name in ['Identificação','Classe','Grupo']:
        if list(roles.values()).count(name) > 1:
            raise ValueError(f'Selecione no máximo uma coluna para {name}.')
    if 'X' not in roles.values():
        raise ValueError('Selecione pelo menos uma variável X.')
    ex = cfg['excluded']
    if not isinstance(ex, list) or any(type(i) is not int or i < 0 or i >= len(df) for i in ex) or len(ex) != len(set(ex)) or len(ex) >= len(df):
        raise ValueError('Seleção de amostras inválida; mantenha ao menos uma.')
    if any(not isinstance(cfg[k],str) or len(cfg[k]) > 10_000 for k in ['axis_label','axis_unit','notes']):
        raise ValueError('Metadados de texto inválidos.')
    idcols = [c for c,v in roles.items() if v == 'Identificação']
    if idcols:
        ids = df[idcols[0]].str.strip()
        if ids.eq('').any() or ids.duplicated().any():
            raise ValueError('Identificações devem ser preenchidas e únicas; use Grupo para réplicas.')
    for c,v in roles.items():
        if v in ['X','Y']:
            nums = numeric_column(df[c], cfg['decimal'])
            invalid = df[c].str.strip().ne('') & nums.isna()
            if invalid.any() or np.isinf(nums.to_numpy(dtype=float)).any():
                raise ValueError(f'Coluna {c}: texto ou infinito em X/Y. Confira o decimal ou altere o papel.')
    if cfg['axis_mode'] == 'Cabeçalhos numéricos':
        x = [c for c,v in roles.items() if v == 'X']
        vals = numeric_column(pd.Series(x), cfg['decimal']).to_numpy(float)
        if not np.isfinite(vals).all() or len(np.unique(vals)) != len(vals):
            raise ValueError('Para usar cabeçalhos como eixo, os nomes de X devem ser números únicos.')
    return cfg


def build_dataset(df, cfg):
    validate_config(df, cfg)
    active = df.drop(index=cfg['excluded'])
    roles = cfg['roles']
    def cols(role): return [c for c,v in roles.items() if v == role]
    def matrix(role):
        return pd.DataFrame({c:numeric_column(active[c],cfg['decimal']) for c in cols(role)}, index=active.index)
    ids = active[cols('Identificação')[0]].tolist() if cols('Identificação') else [f'A{i+1:05d}' for i in active.index]
    x = matrix('X')
    axis = numeric_column(pd.Series(x.columns),cfg['decimal']).to_numpy(float) if cfg['axis_mode'] == 'Cabeçalhos numéricos' else np.arange(1,x.shape[1]+1)
    return {'X':x, 'Y':matrix('Y'), 'ids':ids, 'axis':axis,
            'classes':active[cols('Classe')[0]].tolist() if cols('Classe') else None,
            'groups':active[cols('Grupo')[0]].tolist() if cols('Grupo') else None}


def demo():
    rng = np.random.default_rng(1729)
    axis = np.linspace(900,1700,161)
    c = rng.uniform(0.5,5,36)
    b = rng.uniform(0.3,2,36)
    x = c[:,None]*np.exp(-0.5*((axis-1120)/45)**2) + b[:,None]*np.exp(-0.5*((axis-1480)/70)**2) + rng.normal(0,.025,(36,161))
    df = pd.DataFrame(x,columns=[f'{v:g}' for v in axis])
    df.insert(0,'Concentracao',c)
    df.insert(0,'Grupo',[f'Lote-{i//3+1:02d}' for i in range(36)])
    df.insert(0,'Classe',np.where(c>=2.75,'Alta','Baixa'))
    df.insert(0,'Amostra',[f'SYN-{i+1:03d}' for i in range(36)])
    df = normalize_table(df)
    cfg=default_config(df)
    cfg['roles'].update(Amostra='Identificação',Classe='Classe',Grupo='Grupo',Concentracao='Y')
    cfg.update(axis_mode='Cabeçalhos numéricos',axis_label='Comprimento de onda',axis_unit='nm',notes='Dados sintéticos, semente 1729. Concentração em unidades arbitrárias; sem significado experimental.')
    return df,cfg


def save_project(df, cfg, name):
    validate_config(df,cfg)
    if not isinstance(name,str) or not name.strip() or len(name)>200:
        raise ValueError('Nome do projeto deve ter entre 1 e 200 caracteres.')
    payload={'format':'marvin-chemometrics','schema':1,'app_version':VERSION,'name':name,
             'columns':df.columns.tolist(),'data':df.values.tolist(),'config':cfg}
    data=json.dumps(payload,ensure_ascii=False,allow_nan=False).encode('utf-8')
    if len(data)>MAX_EXPANDED:
        raise ValueError('Projeto excede 64 MB descompactado.')
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('project.json',data)
    raw=out.getvalue()
    if len(raw)>MAX_BYTES:
        raise ValueError('Projeto excede 20 MB compactado.')
    return raw


def load_project(raw):
    try:
        with check_zip(raw,exact=['project.json']) as z:
            p=json.loads(z.read('project.json'))
        if not isinstance(p,dict) or p.get('format')!='marvin-chemometrics' or type(p.get('schema')) is not int or p['schema']!=1:
            raise ValueError('Formato ou versão de projeto não reconhecido.')
        cols=p['columns']; rows=p['data']
        if not isinstance(cols,list) or not cols or not all(isinstance(c,str) for c in cols) or not isinstance(rows,list) or len(cols)>10_000 or len(rows)>20_000 or len(cols)*len(rows)>MAX_CELLS:
            raise ValueError('Dimensões de projeto inválidas.')
        if any(not isinstance(r,list) or len(r)!=len(cols) or any(not isinstance(v,str) for v in r) for r in rows):
            raise ValueError('Tabela de projeto inválida.')
        df=normalize_table(pd.DataFrame(rows,columns=cols))
        if not isinstance(p['name'],str) or not p['name'].strip() or len(p['name'])>200:
            raise ValueError('Nome de projeto inválido.')
        validate_config(df,p['config'])
        return df,p['config'],p['name']
    except (KeyError,TypeError,zipfile.BadZipFile,UnicodeDecodeError,json.JSONDecodeError,RuntimeError) as e:
        raise ValueError('Arquivo de projeto inválido ou corrompido.') from e
